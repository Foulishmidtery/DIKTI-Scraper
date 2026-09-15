"""PostgreSQL persistence and database-backed export/analytics services."""

from __future__ import annotations

import hashlib
import re
import threading
import unicodedata
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any, Iterable

from sqlalchemy import (
    and_,
    BigInteger,
    Column,
    DateTime,
    ForeignKey,
    Integer,
    MetaData,
    String,
    Table,
    Text,
    create_engine,
    func,
    literal_column,
    or_,
    select,
    update,
)
from sqlalchemy.dialects.postgresql import JSONB, insert as pg_insert
from sqlalchemy.engine import Engine
from backend.config import reload_database_settings_if_changed, settings


metadata = MetaData()

scrape_runs = Table(
    "scrape_runs",
    metadata,
    Column("id", String(36), primary_key=True),
    Column("status", String(24), nullable=False),
    Column("selected_prodi", JSONB, nullable=False),
    Column("started_at", DateTime(timezone=True), nullable=False),
    Column("finished_at", DateTime(timezone=True)),
    Column("dosen_seen", Integer, nullable=False, server_default="0"),
    Column("dosen_inserted", Integer, nullable=False, server_default="0"),
    Column("prodi_seen", Integer, nullable=False, server_default="0"),
    Column("prodi_inserted", Integer, nullable=False, server_default="0"),
    Column("export_filename", String(255)),
    Column("error_message", Text),
)

dosen_records = Table(
    "dosen_records",
    metadata,
    Column("id", BigInteger, primary_key=True, autoincrement=True),
    Column("source_key", String(64), nullable=False, unique=True, index=True),
    Column("identity_kind", String(24), nullable=False),
    Column("identity_value", String(255), nullable=False),
    Column("payload", JSONB, nullable=False),
    Column("first_seen_run_id", String(36), ForeignKey("scrape_runs.id", ondelete="SET NULL")),
    Column("first_seen_at", DateTime(timezone=True), nullable=False),
)

prodi_records = Table(
    "prodi_records",
    metadata,
    Column("id", BigInteger, primary_key=True, autoincrement=True),
    Column("source_key", String(64), nullable=False, unique=True, index=True),
    Column("identity_kind", String(24), nullable=False),
    Column("identity_value", String(500), nullable=False),
    Column("payload", JSONB, nullable=False),
    Column("first_seen_run_id", String(36), ForeignKey("scrape_runs.id", ondelete="SET NULL")),
    Column("first_seen_at", DateTime(timezone=True), nullable=False),
)

_engine: Engine | None = None
_engine_signature: tuple[str, int] | None = None
_schema_initialized = False
_schema_lock = threading.Lock()
_engine_lock = threading.Lock()


class DatabaseNotConfigured(RuntimeError):
    pass


def _gateway_client():
    from backend.secure_gateway import gateway_client
    return gateway_client


def _database_url() -> str:
    url = settings.database_url
    if not url:
        raise DatabaseNotConfigured("DATABASE_URL belum dikonfigurasi.")
    if url.startswith("postgres://"):
        url = "postgresql+psycopg://" + url[len("postgres://"):]
    elif url.startswith("postgresql://"):
        url = "postgresql+psycopg://" + url[len("postgresql://"):]
    if not url.startswith("postgresql+psycopg://"):
        raise DatabaseNotConfigured("DATABASE_URL harus menggunakan PostgreSQL.")
    return url


def get_engine() -> Engine:
    global _engine, _engine_signature, _schema_initialized
    if settings.secure_gateway:
        raise DatabaseNotConfigured("Koneksi database langsung dinonaktifkan pada production.")
    reload_database_settings_if_changed()
    current_url = _database_url()
    current_signature = (current_url, settings.database_pool_size)
    if _engine is None or current_signature != _engine_signature:
        with _engine_lock:
            if _engine is None or current_signature != _engine_signature:
                if _engine is not None:
                    _engine.dispose()
                _engine = create_engine(
                    current_url,
                    pool_pre_ping=True,
                    pool_size=settings.database_pool_size,
                    max_overflow=5,
                    connect_args={"connect_timeout": 5},
                )
                _engine_signature = current_signature
                _schema_initialized = False
    return _engine


def initialize_database() -> None:
    global _schema_initialized
    if settings.secure_gateway:
        status = _gateway_client().status()
        if not status["connected"]:
            raise DatabaseNotConfigured(status["message"])
        return
    engine = get_engine()
    if not _schema_initialized:
        with _schema_lock:
            if not _schema_initialized:
                metadata.create_all(engine)
                _schema_initialized = True
    with engine.connect() as connection:
        connection.execute(select(1))


def database_status() -> dict[str, Any]:
    if settings.secure_gateway:
        try:
            return _gateway_client().status()
        except Exception:
            return {"configured": bool(settings.gateway_url), "connected": False, "message": "Secure gateway belum siap"}
    reload_database_settings_if_changed()
    configured = bool(settings.database_url)
    if not configured:
        return {"configured": False, "connected": False, "message": "DATABASE_URL belum diatur"}
    try:
        initialize_database()
        with get_engine().connect() as connection:
            dosen_total = connection.scalar(select(func.count()).select_from(dosen_records)) or 0
            prodi_total = connection.scalar(select(func.count()).select_from(prodi_records)) or 0
        return {
            "configured": True,
            "connected": True,
            "message": "PostgreSQL terhubung",
            "total_dosen": int(dosen_total),
            "total_prodi": int(prodi_total),
        }
    except Exception as exc:
        return {"configured": True, "connected": False, "message": _safe_db_error(exc)}


def _safe_db_error(exc: Exception) -> str:
    text = str(exc).splitlines()[0].strip()
    text = re.sub(r"postgresql[^\s]*://[^@\s]+@", "postgresql://***@", text)
    return text[:240] or exc.__class__.__name__


def _normalize(value: Any) -> str:
    text = unicodedata.normalize("NFKC", str(value or "")).strip().upper()
    return " ".join(text.split())


def _source_key(kind: str, value: str) -> str:
    return hashlib.sha256(f"{kind}|{value}".encode("utf-8")).hexdigest()


def dosen_identity(payload: dict[str, Any]) -> tuple[str, str, str]:
    nidn = _normalize(payload.get("NIDN"))
    if nidn:
        return _source_key("nidn", nidn), "nidn", nidn
    nuptk = _normalize(payload.get("NUPTK"))
    if nuptk:
        return _source_key("nuptk", nuptk), "nuptk", nuptk
    fallback = "|".join([
        _normalize(payload.get("Nama")),
        _normalize(payload.get("Perguruan Tinggi")),
        _normalize(payload.get("Program Studi")),
    ])
    return _source_key("fallback", fallback), "fallback", fallback


def prodi_identity(payload: dict[str, Any]) -> tuple[str, str, str]:
    natural_identity = "|".join([
        _normalize(payload.get("nama")),
        _normalize(payload.get("jenjang")),
        _normalize(payload.get("pt")),
    ])
    # The source key deliberately uses the logical identity so legacy Excel rows
    # (which do not contain the PDDIKTI id) collide with future live scrapes.
    source_key = _source_key("prodi", natural_identity)
    pddikti_id = _normalize(payload.get("id"))
    if pddikti_id:
        return source_key, "pddikti_id", pddikti_id
    return source_key, "logical", natural_identity


_ACCREDITATION_FIELDS = (
    "peringkat_akreditasi_banpt", "nomor_sk_akreditasi", "tanggal_sk_akreditasi",
    "tanggal_akhir_akreditasi", "status_berlaku_sk_akreditasi",
    "sumber_akreditasi", "url_sumber_akreditasi", "url_riwayat_akreditasi", "url_sk_akreditasi",
    "lembaga_akreditasi_nasional", "peringkat_akreditasi_nasional",
)


def _accreditation_date_key(value: Any) -> str:
    raw = str(value or "").strip()
    if re.fullmatch(r"\d{4}-\d{2}-\d{2}", raw):
        return raw
    months = {name: f"{index:02d}" for index, name in enumerate(
        ("januari", "februari", "maret", "april", "mei", "juni", "juli", "agustus",
         "september", "oktober", "november", "desember"), 1)}
    parts = raw.lower().split()
    if len(parts) == 3 and parts[0].isdigit() and parts[1] in months and parts[2].isdigit():
        return f"{int(parts[2]):04d}-{months[parts[1]]}-{int(parts[0]):02d}"
    return _normalize(raw)


def merge_prodi_accreditation(previous: dict[str, Any] | None, incoming: dict[str, Any], run_id: str) -> dict[str, Any]:
    """Preserve issuer evidence and append actual accreditation changes between runs."""
    if not previous:
        return incoming
    merged = dict(incoming)
    old_history = previous.get("riwayat_akreditasi_banpt")
    if not merged.get("riwayat_akreditasi_banpt") and isinstance(old_history, list):
        merged["riwayat_akreditasi_banpt"] = old_history
    old_lam = bool(previous.get("lembaga_akreditasi_nasional"))
    new_lam = merged.get("status_pencocokan_lam") == "terverifikasi"
    verified = new_lam or (not old_lam and merged.get("status_pencocokan_akreditasi") == "terverifikasi")
    old_rank = _normalize(previous.get("peringkat_akreditasi_nasional") or previous.get("peringkat_akreditasi_banpt"))
    new_rank = _normalize(merged.get("peringkat_akreditasi_nasional") or merged.get("peringkat_akreditasi_banpt"))
    same_decision = (old_rank == new_rank
                     and _normalize(previous.get("sumber_akreditasi")) == _normalize(merged.get("sumber_akreditasi"))
                     and _accreditation_date_key(previous.get("tanggal_akhir_akreditasi")) == _accreditation_date_key(merged.get("tanggal_akhir_akreditasi")))
    if verified and same_decision:
        for field in ("nomor_sk_akreditasi", "tanggal_sk_akreditasi", "url_sk_akreditasi"):
            if not merged.get(field) and previous.get(field):
                merged[field] = previous[field]
    if not verified and old_rank:
        for field in _ACCREDITATION_FIELDS:
            if previous.get(field):
                merged[field] = previous[field]
        merged["status_pencocokan_akreditasi"] = "data terakhir terverifikasi; sumber belum sinkron"
    checks = merged.get("pemeriksaan_akreditasi_internasional")
    if (not merged.get("akreditasi_internasional_terverifikasi")
            and previous.get("akreditasi_internasional_terverifikasi")
            and isinstance(checks, list) and checks
            and all(isinstance(item, dict) and item.get("status") == "sumber tidak tersedia" for item in checks)):
        merged["akreditasi_internasional_terverifikasi"] = previous["akreditasi_internasional_terverifikasi"]
        merged["status_pencocokan_internasional"] = "data terakhir terverifikasi; sumber belum tersedia"
    changes = previous.get("riwayat_perubahan_scraping")
    change_rows = list(changes) if isinstance(changes, list) else []
    if verified and old_rank and new_rank and (
        old_rank != new_rank or _accreditation_date_key(previous.get("tanggal_akhir_akreditasi")) != _accreditation_date_key(merged.get("tanggal_akhir_akreditasi"))
    ):
        change_rows.append({
            "run_id": run_id,
            "dicatat_pada": datetime.now(timezone.utc).isoformat(),
            "peringkat_sebelum": previous.get("peringkat_akreditasi_nasional") or previous.get("peringkat_akreditasi_banpt", ""),
            "peringkat_sesudah": merged.get("peringkat_akreditasi_nasional") or merged.get("peringkat_akreditasi_banpt", ""),
            "berlaku_sampai_sebelum": previous.get("tanggal_akhir_akreditasi", ""),
            "berlaku_sampai_sesudah": merged.get("tanggal_akhir_akreditasi", ""),
        })
    if change_rows:
        merged["riwayat_perubahan_scraping"] = change_rows[-100:]
    return merged


def create_scrape_run(run_id: str, selected_prodi: list[str]) -> None:
    if settings.secure_gateway:
        _gateway_client().create_run(run_id, selected_prodi)
        return
    initialize_database()
    with get_engine().begin() as connection:
        connection.execute(pg_insert(scrape_runs).values(
            id=run_id,
            status="running",
            selected_prodi=selected_prodi,
            started_at=datetime.now(timezone.utc),
        ).on_conflict_do_nothing(index_elements=[scrape_runs.c.id]))


def _chunks(rows: list[dict[str, Any]], size: int = 500) -> Iterable[list[dict[str, Any]]]:
    for start in range(0, len(rows), size):
        yield rows[start:start + size]


def persist_scrape_results(run_id: str, profiles: list[dict[str, Any]], prodi_list: list[dict[str, Any]]) -> dict[str, int]:
    if settings.secure_gateway:
        return _gateway_client().persist(run_id, profiles, prodi_list)
    now = datetime.now(timezone.utc)
    dosen_rows = []
    for payload in profiles:
        source_key, identity_kind, identity_value = dosen_identity(payload)
        dosen_rows.append({
            "source_key": source_key,
            "identity_kind": identity_kind,
            "identity_value": identity_value,
            "payload": payload,
            "first_seen_run_id": run_id,
            "first_seen_at": now,
        })

    prodi_rows = []
    for payload in prodi_list:
        source_key, identity_kind, identity_value = prodi_identity(payload)
        prodi_rows.append({
            "source_key": source_key,
            "identity_kind": identity_kind,
            "identity_value": identity_value,
            "payload": payload,
            "first_seen_run_id": run_id,
            "first_seen_at": now,
        })

    dosen_inserted = 0
    dosen_updated = 0
    prodi_inserted = 0
    prodi_updated = 0
    with get_engine().begin() as connection:
        for batch in _chunks(dosen_rows):
            insert_statement = pg_insert(dosen_records).values(batch)
            statement = insert_statement.on_conflict_do_update(
                index_elements=[dosen_records.c.source_key],
                set_={
                    "identity_kind": insert_statement.excluded.identity_kind,
                    "identity_value": insert_statement.excluded.identity_value,
                    "payload": insert_statement.excluded.payload,
                },
                where=and_(
                    insert_statement.excluded.payload["ID Dosen PDDIKTI"].astext != "",
                    dosen_records.c.payload.is_distinct_from(insert_statement.excluded.payload),
                ),
            ).returning(literal_column("(xmax = 0)").label("inserted"))
            states = [bool(row[0]) for row in connection.execute(statement).fetchall()]
            dosen_inserted += sum(states)
            dosen_updated += len(states) - sum(states)
        for batch in _chunks(prodi_rows):
            existing = dict(connection.execute(
                select(prodi_records.c.source_key, prodi_records.c.payload)
                .where(prodi_records.c.source_key.in_([row["source_key"] for row in batch]))
                .order_by(prodi_records.c.source_key).with_for_update()
            ).fetchall())
            for row in batch:
                row["payload"] = merge_prodi_accreditation(existing.get(row["source_key"]), row["payload"], run_id)
            insert_statement = pg_insert(prodi_records).values(batch)
            statement = insert_statement.on_conflict_do_update(
                index_elements=[prodi_records.c.source_key],
                set_={
                    "identity_kind": insert_statement.excluded.identity_kind,
                    "identity_value": insert_statement.excluded.identity_value,
                    "payload": insert_statement.excluded.payload,
                },
                where=and_(
                    insert_statement.excluded.payload["id"].astext != "",
                    prodi_records.c.payload.is_distinct_from(insert_statement.excluded.payload),
                ),
            ).returning(literal_column("(xmax = 0)").label("inserted"))
            states = [bool(row[0]) for row in connection.execute(statement).fetchall()]
            prodi_inserted += sum(states)
            prodi_updated += len(states) - sum(states)

        connection.execute(update(scrape_runs).where(scrape_runs.c.id == run_id).values(
            dosen_seen=len(dosen_rows),
            dosen_inserted=dosen_inserted,
            prodi_seen=len(prodi_rows),
            prodi_inserted=prodi_inserted,
        ))

    return {
        "dosen_seen": len(dosen_rows),
        "dosen_inserted": dosen_inserted,
        "dosen_updated": dosen_updated,
        "dosen_skipped": len(dosen_rows) - dosen_inserted - dosen_updated,
        "prodi_seen": len(prodi_rows),
        "prodi_inserted": prodi_inserted,
        "prodi_updated": prodi_updated,
        "prodi_skipped": len(prodi_rows) - prodi_inserted - prodi_updated,
    }


def complete_scrape_run(run_id: str, filename: str | None = None) -> None:
    if settings.secure_gateway:
        _gateway_client().finish_run(run_id, "done")
        return
    with get_engine().begin() as connection:
        connection.execute(update(scrape_runs).where(scrape_runs.c.id == run_id).values(
            status="done",
            finished_at=datetime.now(timezone.utc),
            export_filename=filename,
        ))


def fail_scrape_run(run_id: str, error: Exception | str) -> None:
    if settings.secure_gateway:
        try:
            _gateway_client().finish_run(run_id, "error", "Proses scraping gagal.")
        except Exception:
            pass
        return
    try:
        with get_engine().begin() as connection:
            connection.execute(update(scrape_runs).where(
                scrape_runs.c.id == run_id,
                scrape_runs.c.status != "cancelled",
            ).values(
                status="error",
                finished_at=datetime.now(timezone.utc),
                error_message=_safe_db_error(error) if isinstance(error, Exception) else str(error)[:500],
            ))
    except Exception:
        pass


def cancel_scrape_run(run_id: str) -> None:
    if settings.secure_gateway:
        _gateway_client().finish_run(run_id, "cancelled")
        return
    with get_engine().begin() as connection:
        connection.execute(update(scrape_runs).where(scrape_runs.c.id == run_id).values(
            status="cancelled",
            finished_at=datetime.now(timezone.utc),
        ))


def list_scrape_runs(limit: int = 250) -> list[dict[str, Any]]:
    if settings.secure_gateway:
        return _gateway_client().list_runs(max(1, min(limit, 500)))
    initialize_database()
    statement = select(scrape_runs).order_by(scrape_runs.c.started_at.desc()).limit(max(1, min(limit, 500)))
    with get_engine().connect() as connection:
        rows = connection.execute(statement).mappings().all()
    result = []
    for row in rows:
        started_at = row["started_at"]
        finished_at = row["finished_at"]
        result.append({
            "id": row["id"],
            "status": row["status"],
            "selected_prodi": list(row["selected_prodi"] or []),
            "started_at": started_at.isoformat() if started_at else None,
            "finished_at": finished_at.isoformat() if finished_at else None,
            "duration_seconds": max(0, int((finished_at - started_at).total_seconds())) if started_at and finished_at else None,
            "dosen_seen": row["dosen_seen"],
            "dosen_inserted": row["dosen_inserted"],
            "prodi_seen": row["prodi_seen"],
            "prodi_inserted": row["prodi_inserted"],
            "export_filename": row["export_filename"],
            "error_message": row["error_message"],
        })
    return result


_DOSEN_HEAVY_FIELDS = {
    "Data Homebase Mentah", "Data Pencarian Mentah", "Profil Mentah",
    "Riwayat Pendidikan", "Riwayat Mengajar", "Sertifikasi", "Status Detail PDDIKTI",
}


def _compact_dosen(payload: dict[str, Any]) -> dict[str, Any]:
    return {key: value for key, value in payload.items() if key not in _DOSEN_HEAVY_FIELDS}


def get_database_records(*, compact_dosen: bool = False) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    if settings.secure_gateway:
        return (
            _gateway_client().records("dosen", compact=compact_dosen),
            _gateway_client().records("prodi"),
        )
    initialize_database()
    with get_engine().connect() as connection:
        profiles = list(connection.scalars(select(dosen_records.c.payload).order_by(dosen_records.c.id)))
        prodi_list = list(connection.scalars(select(prodi_records.c.payload).order_by(prodi_records.c.id)))
    if compact_dosen:
        profiles = [_compact_dosen(payload) for payload in profiles]
    return profiles, prodi_list


def _prodi_analytics_rows(profiles: list[dict[str, Any]], prodi_list: list[dict[str, Any]]) -> list[dict[str, Any]]:
    counts: dict[str, int] = {}
    for profile in profiles:
        key = f"{_normalize(profile.get('Program Studi'))}|{_normalize(profile.get('Perguruan Tinggi'))}"
        counts[key] = counts.get(key, 0) + 1

    accreditation = {"A": "Unggul", "B": "Baik Sekali", "C": "Baik"}
    result = []
    for index, prodi in enumerate(sorted(prodi_list, key=lambda item: (_normalize(item.get("nama")), _normalize(item.get("pt")))), 1):
        key = f"{_normalize(prodi.get('nama'))}|{_normalize(prodi.get('pt'))}"
        akreditasi = prodi.get("akreditasi", "")
        result.append({
            "No": index,
            "ID Prodi PDDIKTI": prodi.get("id", ""),
            "Kode Prodi": prodi.get("kode_prodi", ""),
            "Nama Prodi": prodi.get("nama", ""),
            "Jenjang": prodi.get("jenjang", ""),
            "Perguruan Tinggi": prodi.get("pt", ""),
            "Jumlah Dosen": counts.get(key, 0),
            "Status Prodi": prodi.get("keterangan", ""),
            "Akreditasi Program Studi": accreditation.get(str(akreditasi), akreditasi),
            "Berlaku Sampai Akreditasi": prodi.get("tanggal_akhir_akreditasi", ""),
            "PTN/PTS": prodi.get("ptn_pts", ""),
            "PTKIN/NON PTKIN": prodi.get("ptkin_non", ""),
            "DIKTI/DIKTIS": prodi.get("dikti_diktis", ""),
            "Pembina": prodi.get("pembina", ""),
            "Provinsi": prodi.get("provinsi", ""),
            "Semester Laporan Terakhir": prodi.get("semester_lapor", "Belum Lapor"),
        })
    return result


def get_analytics() -> dict[str, Any]:
    profiles, prodi_list = get_database_records(compact_dosen=True)
    dosen = [{"No": index, **payload} for index, payload in enumerate(profiles, 1)]
    prodi = _prodi_analytics_rows(profiles, prodi_list)
    return {
        "dosen": dosen,
        "prodi": prodi,
        "metadata": {
            "source": "postgresql",
            "generated_at": datetime.now(timezone.utc).isoformat(),
            "total_dosen": len(dosen),
            "total_prodi": len(prodi),
        },
    }


def get_dosen_detail(pddikti_id: str = "", nidn: str = "") -> dict[str, Any] | None:
    """Ambil satu payload lengkap secara on-demand agar halaman analitik tetap ringan."""
    clean_id = str(pddikti_id or "").strip()[:255]
    clean_nidn = _normalize(nidn)[:255]
    if not clean_id and not clean_nidn:
        raise ValueError("Identitas dosen diperlukan.")
    if settings.secure_gateway:
        result = _gateway_client().request("dosen_detail", {
            "pddikti_id": clean_id,
            "nidn": clean_nidn,
        })
        data = result.get("data")
        return data if isinstance(data, dict) else None

    initialize_database()
    conditions = []
    if clean_nidn:
        conditions.append(dosen_records.c.source_key == _source_key("nidn", clean_nidn))
    if clean_id:
        conditions.append(dosen_records.c.payload["ID Dosen PDDIKTI"].astext == clean_id)
    statement = select(dosen_records.c.payload).where(or_(*conditions)).limit(1)
    with get_engine().connect() as connection:
        payload = connection.scalar(statement)
    return payload if isinstance(payload, dict) else None


def get_prodi_detail(pddikti_id: str = "", nama: str = "", jenjang: str = "", pt: str = "") -> dict[str, Any] | None:
    """Ambil satu payload prodi lengkap secara on-demand."""
    clean_id = str(pddikti_id or "").strip()[:255]
    logical_identity = "|".join([_normalize(nama), _normalize(jenjang), _normalize(pt)])
    logical_key = _source_key("prodi", logical_identity) if all([nama, jenjang, pt]) else ""
    if not clean_id and not logical_key:
        raise ValueError("Identitas program studi diperlukan.")
    if settings.secure_gateway:
        result = _gateway_client().request("prodi_detail", {
            "pddikti_id": clean_id,
            "nama": str(nama or "")[:200],
            "jenjang": str(jenjang or "")[:32],
            "pt": str(pt or "")[:255],
        })
        data = result.get("data")
        return data if isinstance(data, dict) else None

    initialize_database()
    conditions = []
    if clean_id:
        conditions.append(prodi_records.c.payload["id"].astext == clean_id)
    if logical_key:
        conditions.append(prodi_records.c.source_key == logical_key)
    statement = select(prodi_records.c.payload).where(or_(*conditions)).limit(1)
    with get_engine().connect() as connection:
        payload = connection.scalar(statement)
    return payload if isinstance(payload, dict) else None


def export_database(output_dir: str, callback=lambda _message: None) -> str:
    from scraper.dosen_scraper import export_to_excel

    profiles, prodi_list = get_database_records()
    if not profiles and not prodi_list:
        raise RuntimeError("Database belum memiliki data untuk diekspor.")
    semesters = sorted({_normalize(row.get("Semester Data")) for row in profiles if row.get("Semester Data")})
    semester_label = ", ".join(semesters[-3:]) if semesters else "Database"
    return export_to_excel(profiles, prodi_list, semester_label, output_dir, callback)


def excel_rows(path: str | Path) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        profiles: list[dict[str, Any]] = []
        prodi_list: list[dict[str, Any]] = []

        def rows(sheet_name: str, header_row: int, first_row: int) -> list[dict[str, Any]]:
            sheet = workbook[sheet_name]
            headers = [cell.value for cell in sheet[header_row]]
            result = []
            for values in sheet.iter_rows(min_row=first_row, values_only=True):
                if not values or not any(value is not None for value in values):
                    continue
                result.append({
                    header: value.isoformat() if isinstance(value, (date, datetime)) else (value if value is not None else "")
                    for header, value in zip(headers, values) if header
                })
            return result

        if "Dosen" in workbook.sheetnames:
            dosen_rows = rows("Dosen", 5, 6)
            for record in dosen_rows:
                profiles.append({
                    "Nama": record.get("Nama", ""),
                    "Perguruan Tinggi": record.get("Perguruan Tinggi", ""),
                    "Program Studi": record.get("Program Studi", ""),
                    "Pendidikan Terakhir": record.get("Pendidikan Terakhir", ""),
                    "Status Aktifitas": record.get("Status Aktivitas", ""),
                    "Status Ikatan Kerja": record.get("Ikatan Kerja", ""),
                    "Jabatan Fungsional": record.get("Jabatan Fungsional", ""),
                    "Jenis Kelamin": record.get("Jenis Kelamin", ""),
                    "Status Kepegawaian": record.get("Status Kepegawaian", ""),
                    "NIDN": record.get("NIDN", ""),
                    "NUPTK": record.get("NUPTK", ""),
                    "Semester Data": record.get("Semester Data", ""),
                    "Riwayat Pendidikan": [],
                    "Riwayat Mengajar": [],
                    "Sertifikasi": [],
                })
            by_number = {
                str(record.get("No Dosen")): profile
                for record, profile in zip(dosen_rows, profiles)
            }
            if "Pendidikan" in workbook.sheetnames:
                for record in rows("Pendidikan", 5, 6):
                    profile = by_number.get(str(record.get("No Dosen")))
                    if profile is not None:
                        profile["Riwayat Pendidikan"].append({
                            "jenjang": record.get("Jenjang Pendidikan", ""),
                            "nama_pt": record.get("PT Pendidikan", ""),
                            "nama_prodi": record.get("Program Studi Pendidikan", ""),
                            "tahun_masuk": record.get("Tahun Masuk", ""),
                            "tahun_lulus": record.get("Tahun Lulus", ""),
                            "gelar": record.get("Gelar", ""),
                            "singkatan_gelar": record.get("Singkatan Gelar", ""),
                        })
            if "Mengajar" in workbook.sheetnames:
                for record in rows("Mengajar", 5, 6):
                    profile = by_number.get(str(record.get("No Dosen")))
                    if profile is not None:
                        profile["Riwayat Mengajar"].append({
                            "nama_pt": record.get("PT Mengajar", ""),
                            "nama_matkul": record.get("Mata Kuliah", ""),
                            "kode_matkul": record.get("Kode Matkul", ""),
                            "nama_kelas": record.get("Kelas", ""),
                            "nama_semester": record.get("Semester", ""),
                        })
            if "Sertifikasi" in workbook.sheetnames:
                for record in rows("Sertifikasi", 5, 6):
                    profile = by_number.get(str(record.get("No Dosen")))
                    if profile is not None:
                        profile["Sertifikasi"].append({
                            key: value for key, value in record.items()
                            if key not in {"No Dosen", "Nama", "PT Dosen"}
                        })
        elif "Data Dosen" in workbook.sheetnames:
            for record in rows("Data Dosen", 5, 6):
                record.pop("No", None)
                profiles.append(record)

        prodi_sheet = "Prodi" if "Prodi" in workbook.sheetnames else "Daftar Prodi"
        if prodi_sheet in workbook.sheetnames:
            for record in rows(prodi_sheet, 5 if prodi_sheet == "Prodi" else 3, 6 if prodi_sheet == "Prodi" else 4):
                prodi_list.append({
                    "kode_prodi": record.get("Kode Prodi", ""),
                    "nama": record.get("Nama Prodi", ""),
                    "jenjang": record.get("Jenjang", ""),
                    "pt": record.get("Perguruan Tinggi", ""),
                    "keterangan": record.get("Status Prodi", record.get("Keterangan", "")),
                    "akreditasi": record.get("Akreditasi Program Studi", ""),
                    "tanggal_akhir_akreditasi": record.get("Berlaku Sampai Akreditasi", ""),
                    "nomor_sk_akreditasi": record.get("Nomor SK Akreditasi", ""),
                    "tanggal_sk_akreditasi": record.get("Tanggal SK Akreditasi", ""),
                    "status_berlaku_sk_akreditasi": record.get("Status Berlaku SK Akreditasi", ""),
                    "ptn_pts": record.get("PTN/PTS", ""),
                    "ptkin_non": record.get("PTKIN/NON PTKIN", ""),
                    "dikti_diktis": record.get("DIKTI/DIKTIS", ""),
                    "pembina": record.get("Pembina", ""),
                    "provinsi": record.get("Provinsi", ""),
                    "tanggal_berdiri": record.get("Tanggal Berdiri", ""),
                    "nomor_sk_penyelenggaraan": record.get("Nomor SK Penyelenggaraan", record.get("Nomor SK Prodi", "")),
                    "tanggal_sk_penyelenggaraan": record.get("Tanggal SK Penyelenggaraan", record.get("Tanggal SK Prodi", "")),
                    "semester_lapor": record.get("Semester Laporan Terakhir", "Belum Lapor"),
                })
            if "Riwayat Akreditasi" in workbook.sheetnames:
                by_prodi = {
                    (str(item.get("nama") or "").upper(), str(item.get("jenjang") or "").upper(),
                     str(item.get("pt") or "").upper()): item
                    for item in prodi_list
                }
                for record in rows("Riwayat Akreditasi", 5 if prodi_sheet == "Prodi" else 1, 6 if prodi_sheet == "Prodi" else 2):
                    key = (
                        str(record.get("Nama Prodi") or "").upper(),
                        str(record.get("Jenjang") or "").upper(),
                        str(record.get("Perguruan Tinggi") or "").upper(),
                    )
                    item = by_prodi.get(key)
                    if item is not None:
                        item.setdefault("riwayat_akreditasi_banpt", []).append({
                            "tanggal_sk": record.get("Tanggal SK", ""),
                            "peringkat": record.get("Peringkat", ""),
                            "berlaku_sampai": record.get("Berlaku Sampai", ""),
                            "status_berlaku_sk": record.get("Status Berlaku SK", ""),
                            "sumber": record.get("Sumber", ""),
                        })
        return profiles, prodi_list
    finally:
        workbook.close()
