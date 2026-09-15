"""Database/scraper-backed Excel export with readable, filterable data tables."""

import json
import os
import re
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


MONTHS = {
    "januari": 1, "februari": 2, "maret": 3, "april": 4, "mei": 5,
    "juni": 6, "juli": 7, "agustus": 8, "september": 9, "oktober": 10,
    "november": 11, "desember": 12,
}
PALETTE = {
    "Dosen": ("087F5B", "E8F7EF", "TableStyleMedium4"),
    "Prodi": ("2869BD", "EDF5FF", "TableStyleMedium2"),
    "Pendidikan": ("7653A7", "F3EDFA", "TableStyleMedium5"),
    "Mengajar": ("A97013", "FFF6E4", "TableStyleMedium9"),
    "Riwayat Akreditasi": ("0F766E", "E5F6F2", "TableStyleMedium4"),
    "Sertifikasi": ("9A5B15", "FFF4DF", "TableStyleMedium9"),
}


def _cell_value(value):
    if value is None:
        return None
    if isinstance(value, (date, datetime, int, float, bool)):
        return value
    if isinstance(value, (dict, list)):
        value = json.dumps(value, ensure_ascii=False, separators=(",", ":"))
    text = str(value)
    if text.startswith(("=", "+", "-", "@")):
        text = "'" + text
    return text[:32767]


def _date_value(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value or "").strip()
    if not text:
        return None
    match = re.match(r"^(\d{4})-(\d{2})-(\d{2})", text)
    if match:
        try:
            return date(*map(int, match.groups()))
        except ValueError:
            return _cell_value(text)
    match = re.match(r"^(\d{1,2})\s+([A-Za-z]+)\s+(\d{4})$", text)
    if match and match.group(2).lower() in MONTHS:
        try:
            return date(int(match.group(3)), MONTHS[match.group(2).lower()], int(match.group(1)))
        except ValueError:
            pass
    return _cell_value(text)


def _pick(record, *keys):
    for key in keys:
        value = record.get(key)
        if value is not None and value != "":
            return value
    return None


def _year(value):
    text = str(value or "").strip()
    return int(text) if re.fullmatch(r"\d{4}", text) else (text or None)


def _sheet(workbook, title, headers, rows, widths, *, date_columns=(), text_columns=()):
    accent, soft, table_style = PALETTE[title]
    ws = workbook.create_sheet(title)
    ws.sheet_properties.tabColor = accent
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3
    ws["B2"] = title if title != "Prodi" else "Daftar program studi"
    ws["B2"].font = Font(name="Arial", size=15, bold=True, color="123B34")
    ws["B3"] = f"{len(rows):,} baris"
    ws["B3"].font = Font(name="Arial", size=10, italic=True, color="5C706A")
    ws["B4"] = "Filter dan urutkan melalui panah pada setiap judul kolom"
    ws["B4"].font = Font(name="Arial", size=10, bold=True, color="123B34")
    ws["B4"].fill = PatternFill("solid", fgColor=soft)
    for index, heading in enumerate(headers, 2):
        cell = ws.cell(5, index, heading)
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=accent)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[5].height = 28
    for index, width in enumerate(widths, 2):
        ws.column_dimensions[get_column_letter(index)].width = width
    for row_number, record in enumerate(rows, 6):
        for col_number, value in enumerate(record, 2):
            cell = ws.cell(row_number, col_number)
            if col_number - 2 in date_columns:
                cell.value = _date_value(value)
            elif col_number - 2 in text_columns and value is not None:
                cell.value = _cell_value(str(value))
            else:
                cell.value = _cell_value(value)
            if col_number - 2 in date_columns and isinstance(cell.value, date):
                cell.number_format = "dd mmm yyyy"
            elif col_number - 2 in text_columns:
                cell.number_format = "@"
            elif isinstance(cell.value, (int, float)):
                cell.number_format = "#,##0"
    if rows:
        table = Table(displayName="Tabel" + title.replace(" ", ""), ref=f"B5:{get_column_letter(len(headers) + 1)}{5 + len(rows)}")
        table.tableStyleInfo = TableStyleInfo(name=table_style, showRowStripes=True)
        ws.add_table(table)
    ws.freeze_panes = "C6"
    return ws


def _highlight(ws, column, last_row, formula, fill, color):
    if last_row < 6:
        return
    ws.conditional_formatting.add(
        f"{column}6:{column}{last_row}",
        FormulaRule(formula=[formula], fill=PatternFill("solid", fgColor=fill),
                    font=Font(name="Arial", color=color)),
    )


def export_to_excel(profiles, prodi_list, semester, output_dir, callback):
    callback("\nSTEP 4: Membuat file Excel...")
    callback({"__progress__": True, "step": 4, "current": 0, "total": 1, "label": "Export Excel"})
    workbook = Workbook()
    workbook.remove(workbook.active)
    sorted_profiles = sorted(profiles, key=lambda item: (str(item.get("Program Studi") or ""), str(item.get("Nama") or "")))
    sorted_prodi = sorted(prodi_list, key=lambda item: (str(item.get("nama") or ""), str(item.get("pt") or "")))
    dosen_number = {id(profile): index for index, profile in enumerate(sorted_profiles, 1)}

    dosen_headers = ["No Dosen", "Nama", "Perguruan Tinggi", "Program Studi", "Pendidikan Terakhir",
                     "Status Aktivitas", "Ikatan Kerja", "Jabatan Fungsional", "Jenis Kelamin",
                     "Status Kepegawaian", "NIDN", "NUPTK", "Semester Data",
                     "Riwayat Pendidikan", "Riwayat Mengajar"]
    dosen_rows = [
        [dosen_number[id(p)], _pick(p, "Nama"), _pick(p, "Perguruan Tinggi"),
         _pick(p, "Program Studi"), _pick(p, "Pendidikan Terakhir"),
         _pick(p, "Status Aktifitas", "Status Aktivitas"), _pick(p, "Status Ikatan Kerja"),
         _pick(p, "Jabatan Fungsional"), _pick(p, "Jenis Kelamin"),
         _pick(p, "Status Kepegawaian"), _pick(p, "NIDN"), _pick(p, "NUPTK"),
         _pick(p, "Semester Data"), len(p.get("Riwayat Pendidikan") or []),
         len(p.get("Riwayat Mengajar") or [])]
        for p in sorted_profiles
    ]
    dosen_ws = _sheet(workbook, "Dosen", dosen_headers, dosen_rows,
                      [11, 27, 39, 28, 19, 19, 17, 19, 16, 20, 17, 20, 14, 19, 17],
                      text_columns=(10, 11))
    _highlight(dosen_ws, "G", len(dosen_rows) + 5, 'G6="Aktif"', "DCF6E8", "08684C")
    _highlight(dosen_ws, "G", len(dosen_rows) + 5, 'ISNUMBER(SEARCH("BELAJAR",G6))', "FFF0CF", "99631C")

    counts = {}
    for p in sorted_profiles:
        key = (str(p.get("Program Studi") or "").strip().upper(), str(p.get("Perguruan Tinggi") or "").strip().upper())
        counts[key] = counts.get(key, 0) + 1
    rank = {"A": "Unggul", "B": "Baik Sekali", "C": "Baik"}
    prodi_headers = ["No", "Kode Prodi", "Nama Prodi", "Jenjang", "Perguruan Tinggi",
                     "Jumlah Dosen", "Status Prodi", "Akreditasi Program Studi",
                     "Berlaku Sampai Akreditasi", "Tanggal Berdiri",
                     "Nomor SK Penyelenggaraan", "Tanggal SK Penyelenggaraan",
                     "Provinsi", "PTN/PTS", "DIKTI/DIKTIS", "Semester Laporan Terakhir",
                     "Lembaga Akreditasi Nasional", "Peringkat Akreditasi Nasional",
                     "Nomor SK Akreditasi", "Tanggal SK Akreditasi",
                     "Status Berlaku SK Akreditasi", "PTKIN/NON PTKIN", "Pembina"]
    prodi_rows = []
    for index, p in enumerate(sorted_prodi, 1):
        key = (str(p.get("nama") or "").strip().upper(), str(p.get("pt") or "").strip().upper())
        raw_rank = p.get("akreditasi") or ""
        prodi_rows.append([
            index, p.get("kode_prodi"), p.get("nama"), p.get("jenjang"), p.get("pt"),
            counts.get(key, 0), p.get("keterangan"), rank.get(raw_rank, raw_rank),
            p.get("tanggal_akhir_akreditasi"), p.get("tanggal_berdiri"),
            p.get("nomor_sk_penyelenggaraan"), p.get("tanggal_sk_penyelenggaraan"),
            p.get("provinsi"), p.get("ptn_pts"), p.get("dikti_diktis"),
            p.get("semester_lapor") or "Belum Lapor",
            p.get("lembaga_akreditasi_nasional") or p.get("sumber_akreditasi"),
            p.get("peringkat_akreditasi_nasional") or p.get("peringkat_akreditasi_banpt"),
            p.get("nomor_sk_akreditasi"),
            p.get("tanggal_sk_akreditasi"), p.get("status_berlaku_sk_akreditasi"),
            p.get("ptkin_non"), p.get("pembina"),
        ])
    prodi_ws = _sheet(workbook, "Prodi", prodi_headers, prodi_rows,
                      [8, 14, 27, 11, 42, 15, 15, 22, 22, 20, 32, 22,
                       22, 13, 17, 22, 24, 24, 32, 22, 25, 20, 28],
                      date_columns=(8, 9, 11, 19), text_columns=(1,))
    _highlight(prodi_ws, "H", len(prodi_rows) + 5, 'H6="Aktif"', "DCF6E8", "08684C")
    _highlight(prodi_ws, "I", len(prodi_rows) + 5, 'I6<>""', "E7F0FF", "205AAB")
    _highlight(prodi_ws, "J", len(prodi_rows) + 5, 'J6=""', "FFF5E8", "A97013")

    education_headers = ["No Dosen", "Nama", "PT Dosen", "Jenjang Pendidikan", "PT Pendidikan",
                         "Program Studi Pendidikan", "Tahun Masuk", "Tahun Lulus",
                         "Gelar", "Singkatan Gelar"]
    education_rows = []
    teaching_headers = ["No Dosen", "Nama", "PT Dosen", "Prodi Dosen",
                        "PT Mengajar", "Mata Kuliah", "Kode Matkul", "Kelas", "Semester"]
    teaching_rows = []
    certification_rows = []
    for p in sorted_profiles:
        identity = [dosen_number[id(p)], p.get("Nama"), p.get("Perguruan Tinggi")]
        for item in p.get("Riwayat Pendidikan") or []:
            if not isinstance(item, dict):
                continue
            education_rows.append(identity + [
                _pick(item, "jenjang", "jenjang_pendidikan"), _pick(item, "nama_pt", "perguruan_tinggi"),
                _pick(item, "nama_prodi", "program_studi"),
                _year(_pick(item, "tahun_masuk")), _year(_pick(item, "tahun_lulus")),
                _pick(item, "gelar"), _pick(item, "singkatan_gelar"),
            ])
        for item in p.get("Riwayat Mengajar") or []:
            if not isinstance(item, dict):
                continue
            teaching_rows.append(identity + [
                p.get("Program Studi"), _pick(item, "nama_pt", "perguruan_tinggi"),
                _pick(item, "nama_matkul", "nama_mata_kuliah"),
                _pick(item, "kode_matkul", "kode_mata_kuliah"),
                _pick(item, "nama_kelas"), _pick(item, "nama_semester", "semester"),
            ])
        for item in p.get("Sertifikasi") or []:
            if isinstance(item, dict):
                certification_rows.append((identity, item))
    _sheet(workbook, "Pendidikan", education_headers, education_rows,
           [11, 28, 40, 19, 39, 32, 15, 15, 27, 20])
    _sheet(workbook, "Mengajar", teaching_headers, teaching_rows,
           [11, 28, 39, 26, 39, 35, 17, 17, 22], text_columns=(6,))

    history_headers = ["No", "Nama Prodi", "Jenjang", "Perguruan Tinggi", "Tanggal SK",
                       "Peringkat", "Berlaku Sampai", "Status Berlaku SK", "Sumber"]
    history_rows = []
    for p in sorted_prodi:
        for item in p.get("riwayat_akreditasi_banpt") or []:
            if isinstance(item, dict):
                history_rows.append([
                    len(history_rows) + 1, p.get("nama"), p.get("jenjang"), p.get("pt"),
                    item.get("tanggal_sk"), item.get("peringkat"), item.get("berlaku_sampai"),
                    item.get("status_berlaku_sk"), item.get("sumber"),
                ])
    if history_rows:
        _sheet(workbook, "Riwayat Akreditasi", history_headers, history_rows,
               [8, 27, 11, 42, 22, 18, 22, 23, 38], date_columns=(4, 6))

    if certification_rows:
        keys = sorted({str(key) for _, item in certification_rows for key in item})
        rows = [identity + [item.get(key) for key in keys] for identity, item in certification_rows]
        _sheet(workbook, "Sertifikasi", ["No Dosen", "Nama", "PT Dosen"] + keys, rows,
               [11, 28, 40] + [24] * len(keys))

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"data_dosen_{timestamp}.xlsx"
    os.makedirs(output_dir, exist_ok=True)
    workbook.save(os.path.join(output_dir, filename))
    callback(f"✅ File Excel disimpan: {filename}")
    return filename
